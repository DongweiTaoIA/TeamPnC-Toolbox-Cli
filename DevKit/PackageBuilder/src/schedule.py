from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import date, datetime
from collections.abc import Collection

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from .config import Settings
from .business_days import add_business_days

_DATE_FORMATS = ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%d-%b-%Y", "%b %d, %Y")


@dataclass
class MigrationEntry:
    row: int
    sprint: str
    code_freeze: date | None
    prod_date: date
    builder: str
    builder_backup: str
    deployer: str
    deployer_backup: str

    @property
    def key(self) -> str:
        return f"{self.sprint}|{self.prod_date.isoformat()}"


def pick_current_and_next(
    entries: list[MigrationEntry], today: date, days_after: int, holidays: Collection[date] = ()
) -> list[MigrationEntry]:
    """Current = earliest migration whose last step is not past; next = the one after it."""
    active = sorted(
        (e for e in entries if add_business_days(e.prod_date, days_after, holidays) >= today),
        key=lambda e: e.prod_date,
    )
    return active[:2]


def _to_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in _DATE_FORMATS:
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    return None


def _to_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def read_schedule(path: str, settings: Settings) -> list[MigrationEntry]:
    """Read migration rows; rows without a parseable PROD date are skipped."""
    indexes = {
        name: column_index_from_string(letter) - 1
        for name, letter in asdict(settings.columns).items()
    }
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if settings.sheet_name not in workbook.sheetnames:
            raise ValueError(f"Sheet '{settings.sheet_name}' not found in {path}")
        sheet = workbook[settings.sheet_name]

        entries: list[MigrationEntry] = []
        rows = sheet.iter_rows(min_row=settings.first_data_row, values_only=True)
        for row_number, row in enumerate(rows, start=settings.first_data_row):
            def cell(name: str):
                index = indexes[name]
                return row[index] if index < len(row) else None

            prod_date = _to_date(cell("prod_date"))
            if prod_date is None:
                continue
            entries.append(
                MigrationEntry(
                    row=row_number,
                    sprint=_to_text(cell("sprint")),
                    code_freeze=_to_date(cell("code_freeze")),
                    prod_date=prod_date,
                    builder=_to_text(cell("builder")),
                    builder_backup=_to_text(cell("builder_backup")),
                    deployer=_to_text(cell("deployer")),
                    deployer_backup=_to_text(cell("deployer_backup")),
                )
            )
        return entries
    finally:
        workbook.close()
